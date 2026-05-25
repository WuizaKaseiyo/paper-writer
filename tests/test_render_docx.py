"""Tests for render_docx — skipped if python-docx is unavailable."""

from __future__ import annotations

import os
from pathlib import Path

import pytest

# Skip the whole module if python-docx isn't installed.
docx = pytest.importorskip("docx")

from tools.render_docx.render_docx import render_docx


def test_minimal_renders(tmp_path):
    """Smoke test: minimal valid call produces a non-empty .docx file."""
    output = tmp_path / "paper.docx"
    result = render_docx.invoke({
        "title": "A Minimal Test Paper",
        "authors": "Alice Author",
        "abstract": "We test that render_docx produces a non-empty .docx.",
        "sections": {
            "1. Introduction": "First paragraph.\n\nSecond paragraph.",
            "2. Conclusion": "Closing paragraph.",
        },
        "references": "Smith J. (2024). A reference. Journal of Things, 1(1), 1.",
        "output_path": str(output),
    })

    assert result["status"] == "ok"
    assert result["section_count"] == 2
    assert output.exists()
    assert output.stat().st_size > 0

    # Round-trip: re-open with python-docx and check structure
    from docx import Document

    doc = Document(str(output))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "A Minimal Test Paper" in text
    assert "Alice Author" in text
    assert "Abstract." in text
    assert "1. Introduction" in text
    assert "2. Conclusion" in text
    assert "References" in text


def test_sections_with_table(tmp_path):
    """Markdown pipe-table in a section body is rendered as a real Word table."""
    output = tmp_path / "with_table.docx"
    table_md = (
        "| Method | Metric A | Metric B |\n"
        "|--------|----------|----------|\n"
        "| Baseline | 65.4 | 0.42 |\n"
        "| Ours | 72.3 | 0.31 |"
    )
    result = render_docx.invoke({
        "title": "Table Test",
        "sections": {
            "1. Results": f"Headline numbers shown below.\n\n{table_md}\n\nFurther discussion of the table follows.",
        },
        "output_path": str(output),
    })
    assert result["status"] == "ok"

    from docx import Document

    doc = Document(str(output))
    tables = doc.tables
    assert len(tables) == 1, f"expected 1 table, got {len(tables)}"
    assert len(tables[0].rows) == 3  # header + 2 data
    assert tables[0].rows[0].cells[0].text == "Method"
    assert tables[0].rows[2].cells[0].text == "Ours"


def test_display_equation_preserved(tmp_path):
    """$$ ... $$ blocks become centered italic paragraphs."""
    output = tmp_path / "eq.docx"
    result = render_docx.invoke({
        "title": "Equation Test",
        "sections": {
            "1. Method": (
                "We define the loss as\n\n"
                "$$L = \\sum_i (y_i - \\hat{y}_i)^2$$\n\n"
                "Subject to the constraint that weights are non-negative."
            ),
        },
        "output_path": str(output),
    })
    assert result["status"] == "ok"

    from docx import Document

    doc = Document(str(output))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "L = \\sum_i" in text


def test_no_python_docx_returns_clear_error(tmp_path, monkeypatch):
    """If python-docx import fails, the tool returns a clear actionable error."""
    import builtins

    original_import = builtins.__import__

    def blocked_import(name, *args, **kwargs):
        if name == "docx" or name.startswith("docx."):
            raise ImportError("simulated missing python-docx")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", blocked_import)

    result = render_docx.invoke({
        "title": "x",
        "sections": {"1.": "x"},
        "output_path": str(tmp_path / "x.docx"),
    })
    assert result["status"] == "error"
    assert "python-docx is not installed" in result["error"]
    assert "pip install python-docx" in result["error"]
